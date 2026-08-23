import csv
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.models import TableRow, TableSource
from app.services.parsers import sanitize_identifier

MAX_COLUMNS = 60


def parse_spreadsheet(path: Path) -> list[dict[str, Any]]:
    """Return one entry per sheet: {name, columns, rows}."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [_parse_csv(path, path.stem)]
    return _parse_xlsx(path)


def _parse_csv(path: Path, name: str) -> dict[str, Any]:
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        reader = csv.reader(handle)
        rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return {"name": name, "columns": [], "rows": []}
    header = rows[0]
    columns = [sanitize_identifier(h) or f"col_{i}" for i, h in enumerate(header)]
    data = []
    for raw in rows[1:]:
        record = {columns[i]: _coerce(raw[i]) for i in range(min(len(columns), len(raw)))}
        data.append(record)
    return {"name": name, "columns": columns, "rows": data}


def _parse_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            rows_iter = worksheet.iter_rows(values_only=True)
            header_row: tuple | None = None
            data_rows: list[list[Any]] = []

            for values in rows_iter:
                cells = ["" if v is None else v for v in values]
                if not any(str(c).strip() for c in cells):
                    continue
                if header_row is None:
                    header_row = tuple(cells)
                    continue
                data_rows.append(cells)

            if header_row is None:
                continue

            columns = [
                sanitize_identifier(str(h)) or f"col_{i}"
                for i, h in enumerate(header_row[:MAX_COLUMNS])
            ]
            records = [
                {columns[i]: _coerce(row[i] if i < len(row) else None) for i in range(len(columns))}
                for row in data_rows
            ]
            records = [r for r in records if any(v != "" and v is not None for v in r.values())]
            sheets.append({"name": worksheet.title, "columns": columns, "rows": records})
    finally:
        workbook.close()
    return sheets


def _coerce(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float, bool, date, datetime)):
        if isinstance(value, datetime):
            return value.isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return value
    text = str(value).strip()
    if text == "":
        return ""
    cleaned = text.replace(",", "").replace("%", "")
    try:
        number = float(cleaned)
        return int(number) if "." not in cleaned else number
    except ValueError:
        return text


def infer_sql_type(values: list[Any]) -> str:
    """Infer a Postgres type for a column from up to 100 sampled values."""
    sample = [v for v in values[:100] if v != "" and v is not None]
    if not sample:
        return "text"
    types = set()
    for value in sample:
        if isinstance(value, bool):
            types.add("boolean")
        elif isinstance(value, int):
            types.add("bigint")
        elif isinstance(value, float):
            types.add("numeric")
        elif isinstance(value, str):
            if _looks_like_iso_date(value):
                types.add("date")
            elif _is_number(value):
                types.add("numeric")
            else:
                types.add("text")
    if not types:
        return "text"
    if types == {"bigint", "numeric"}:
        return "numeric"
    if len(types) == 1:
        return types.pop()
    return "text"


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


def _looks_like_iso_date(text: str) -> bool:
    return len(text) >= 8 and text[:1].isdigit() and ("-" in text[:5] or "/" in text[:5])


def store_table(
    db: Session,
    tenant_id: uuid.UUID,
    document_id: uuid.UUID | None,
    table_name: str,
    sheet_name: str,
    columns: list[str],
    rows: list[dict[str, Any]],
) -> TableSource:
    source = TableSource(
        tenant_id=tenant_id,
        document_id=document_id,
        name=table_name,
        sheet_name=sheet_name,
        columns=[
            {"name": col, "type": infer_sql_type([row.get(col) for row in rows])} for col in columns
        ],
        row_count=len(rows),
    )
    db.add(source)
    db.flush()

    batch: list[TableRow] = []
    for row_number, row in enumerate(rows, start=1):
        batch.append(
            TableRow(
                tenant_id=tenant_id, table_source_id=source.id, row_number=row_number, data=row
            )
        )
        if len(batch) >= 500:
            db.add_all(batch)
            batch = []
    if batch:
        db.add_all(batch)
    db.flush()
    return source


def build_table_summary_chunks(source: TableSource) -> list[str]:
    """Text chunks describing the table so hybrid search can discover it."""
    schema = ", ".join(f"{c['name']} ({c['type']})" for c in source.columns)
    lines = [
        f"Data table '{source.name}'"
        + (f" from sheet '{source.sheet_name}'" if source.sheet_name else "")
        + f" with {source.row_count} rows.",
        f"Columns: {schema}.",
    ]
    return ["\n".join(lines)]


def fetch_rows_json(db: Session, tenant_id: uuid.UUID, source_id: uuid.UUID, limit: int) -> str:
    import json

    from sqlalchemy import text as sql_text

    payload = db.execute(
        sql_text(
            """
            SELECT COALESCE(jsonb_agg(data), '[]'::jsonb) FROM (
                SELECT data FROM table_rows
                WHERE tenant_id = :tid AND table_source_id = :sid
                ORDER BY row_number LIMIT :lim
            ) sub
            """
        ),
        {"tid": str(tenant_id), "sid": str(source_id), "lim": limit},
    ).scalar()
    return json.dumps(json.loads(payload), ensure_ascii=False)
